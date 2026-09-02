from __future__ import annotations
import io
import zipfile
from pathlib import Path
from echooo.service import Problem

MAX_UPLOAD = 5 * 1024 * 1024


def extract_file(filename: str, data: bytes) -> str:
    if len(data) > MAX_UPLOAD:
        raise Problem("Files must be no larger than 5 MB.")
    suffix = Path(filename).suffix.lower()
    try:
        if suffix in {".txt", ".md", ".csv", ".json"}:
            content = data.decode("utf-8-sig")
        elif suffix == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted or len(reader.pages) > 100:
                raise Problem("Upload an unencrypted PDF with no more than 100 pages.")
            content = "\n\n".join(page.extract_text() or "" for page in reader.pages)
        elif suffix == ".docx":
            from docx import Document
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if sum(i.file_size for i in archive.infolist()) > 20 * 1024 * 1024:
                    raise Problem("The uncompressed file is too large.")
            doc = Document(io.BytesIO(data))
            content = "\n".join([p.text for p in doc.paragraphs] +
                [" | ".join(cell.text for cell in row.cells) for t in doc.tables for row in t.rows])
        else:
            raise Problem("Supported formats: TXT, Markdown, CSV, JSON, PDF, and DOCX.")
    except Problem:
        raise
    except Exception as exc:
        raise Problem("Unable to read this file. Check its format or paste the text instead.") from exc
    if not content.strip():
        raise Problem("No text could be extracted. Run OCR on scanned PDFs first.")
    if len(content) > 100000:
        raise Problem("Extracted text exceeds 100,000 characters. Please split the file.")
    return content.strip()
